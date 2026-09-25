# excel_export.py
# Builds the full model as an .xlsx workbook with live Excel formulas.
#
# Every calculation refers back to the Inputs tab, so an analyst can audit the
# model and change assumptions in Excel. Only the Sensitivities tab holds fixed
# values (Python can't create Excel data tables).
#
# While building, formulas refer to cells by name, e.g. "=[acq_ni]*(1+[acq_ni_g@0])".
# [name] is a single cell and [name@i] is year i+1 of a row. The names are
# swapped for real cell addresses once every tab has been laid out.

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from model.engine import apply_simple_mode
from model.inputs import YEARS

UNITS = {
    "USD": {"amount": 1e6, "amount_label": "$M", "shares": 1e6, "shares_label": "M shares", "price_label": "$"},
    "INR": {"amount": 1e7, "amount_label": "₹ crore", "shares": 1e7, "shares_label": "crore shares", "price_label": "₹"},
}

FORMATS = {
    "amount": '#,##0.0;(#,##0.0);"–"',
    "price": '#,##0.00;(#,##0.00)',
    "eps": '#,##0.000;(#,##0.000)',
    "shares": '#,##0.000;(#,##0.000)',
    "pct": '0.0%;(0.0%);"–"',
    "multiple": '0.0"x";(0.0"x")',
    "number": '0.00',
    "years": '0',
    "text": '@',
}

INPUT_FILL = PatternFill("solid", fgColor="FFF7E6")
HEADING_FONT = Font(bold=True, size=12)
TITLE_FONT = Font(bold=True, size=14)
FIRST_COL, VALUE_COL = 3, 3  # values live in column C; years in C..G
TOKEN = re.compile(r"\[([a-z0-9_]+)(?:@(\d))?\]")


class Book:
    def __init__(self, currency):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.units = UNITS[currency]
        self.refs = {}

    def sheet(self, title):
        return Sheet(self, self.wb.create_sheet(title))

    def resolve(self):
        """Swap every [name] token for its cell address."""
        def address(match):
            ref = self.refs[match.group(1)]
            return ref if match.group(2) is None else ref[int(match.group(2))]

        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.value = TOKEN.sub(address, cell.value)


class Sheet:
    def __init__(self, book, ws):
        self.book, self.ws, self.row = book, ws, 1
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 13
        for col in range(FIRST_COL, FIRST_COL + YEARS):
            ws.column_dimensions[get_column_letter(col)].width = 14

    def address(self, col, row):
        return f"'{self.ws.title}'!${get_column_letter(col)}${row}"

    def unit_label(self, kind):
        units = self.book.units
        return {"amount": units["amount_label"], "price": units["price_label"], "eps": units["price_label"],
                "shares": units["shares_label"], "pct": "%", "multiple": "x", "years": "years"}.get(kind, "")

    def title(self, text, note=None):
        self.ws.cell(self.row, 1, text).font = TITLE_FONT
        self.row += 1
        if note:
            self.note(note)
        self.row += 1

    def heading(self, text, years=False):
        self.row += 1 if self.row > 3 else 0
        self.ws.cell(self.row, 1, text).font = HEADING_FONT
        if years:
            for i in range(YEARS):
                cell = self.ws.cell(self.row, FIRST_COL + i, f"Year {i + 1}")
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="right")
        self.row += 1

    def note(self, text):
        cell = self.ws.cell(self.row, 1, text)
        cell.font = Font(italic=True, color="666666")
        self.row += 1

    def _write(self, col, value, kind, is_input):
        cell = self.ws.cell(self.row, col, value)
        cell.number_format = FORMATS.get(kind, "General")
        if kind != "text":
            cell.alignment = Alignment(horizontal="right")
        if is_input:
            cell.fill = INPUT_FILL
        return cell

    def register(self, key, ref):
        if key in self.book.refs:
            raise ValueError(f"Duplicate cell name: {key}")
        self.book.refs[key] = ref

    def value(self, key, label, value, kind="amount", is_input=False):
        self.ws.cell(self.row, 1, label)
        self.ws.cell(self.row, 2, self.unit_label(kind))
        self._write(VALUE_COL, value, kind, is_input)
        if key:
            self.register(key, self.address(VALUE_COL, self.row))
        self.row += 1

    def years(self, key, label, values, kind="amount", is_input=False, bold=False):
        """`values` is a list of 5 values, or a function of the year index i (0-4)."""
        self.ws.cell(self.row, 1, label).font = Font(bold=bold)
        self.ws.cell(self.row, 2, self.unit_label(kind))
        items = values if isinstance(values, list) else [values(i) for i in range(YEARS)]
        for i, v in enumerate(items):
            cell = self._write(FIRST_COL + i, v, kind, is_input)
            cell.font = Font(bold=bold)
        if key:
            self.register(key, [self.address(FIRST_COL + i, self.row) for i in range(YEARS)])
        self.row += 1


def straight_line(amount, life, i):
    """Excel version of ppa.straight_line for year i+1."""
    return f"{amount}/{life}*MIN(MAX({life}-{i},0),1)"


# ---------------------------------------------------------------- Inputs

def write_inputs(book, d):
    s = book.sheet("Inputs")
    amt, sh = book.units["amount"], book.units["shares"]
    a, t = d.acquirer, d.target
    s.title("Inputs", "Shaded cells are assumptions. Every other tab is calculated from them with live formulas.")
    if d.mode == "simple":
        s.note("Simple mode: the payment mix was converted to cash and debt amounts; advanced features are off.")

    s.heading("General")
    s.value("currency", "Currency", d.currency, "text", True)
    s.value("tax", "Tax rate", d.tax_rate, "pct", True)

    def company(prefix, c, name):
        s.heading(name, years=False)
        s.value(f"{prefix}_ni", "Net income (base year)", c.net_income / amt, "amount", True)
        s.value(f"{prefix}_shares", "Diluted shares", c.diluted_shares / sh, "shares", True)
        s.value(f"{prefix}_rev", "Revenue (base year)", c.revenue / amt, "amount", True)
        s.value(f"{prefix}_ebitda", "EBITDA (base year)", c.ebitda / amt, "amount", True)
        s.heading(f"{name}: forecast assumptions", years=True)
        s.years(f"{prefix}_rev_g", "Revenue growth", c.revenue_growth, "pct", True)
        s.years(f"{prefix}_ni_g", "Net income growth", c.net_income_growth, "pct", True)
        if c.ebitda_margin is None:
            s.years(f"{prefix}_margin", "EBITDA margin (base-year margin)",
                    lambda i: f"=IF([{prefix}_rev]>0,[{prefix}_ebitda]/[{prefix}_rev],0)", "pct")
        else:
            s.years(f"{prefix}_margin", "EBITDA margin", c.ebitda_margin, "pct", True)

    company("acq", a, "Acquirer")
    s.heading("Acquirer: price, debt and cash")
    s.value("acq_price", "Share price", a.share_price, "price", True)
    s.value("acq_debt", "Existing debt", a.debt / amt, "amount", True)
    s.value("acq_int", "Existing interest expense", a.interest_expense / amt, "amount", True)
    s.value("acq_cash_flag", "Cash balance entered? (1 = yes)", 0 if a.cash is None else 1, "number", True)
    s.value("acq_cash", "Cash balance", (a.cash or 0) / amt, "amount", True)
    s.value("acq_min_cash", "Minimum cash balance", a.min_cash / amt, "amount", True)
    s.value("acq_cash_rate", "Interest earned on cash", a.interest_rate_on_cash, "pct", True)

    company("tgt", t, "Target")
    s.heading("Target: price, balance sheet")
    s.value("tgt_unaff", "Unaffected share price", t.unaffected_price, "price", True)
    s.value("tgt_debt", "Debt", t.debt / amt, "amount", True)
    s.value("tgt_cash", "Cash", t.cash / amt, "amount", True)
    s.value("tgt_pref", "Preferred stock", t.preferred_stock / amt, "amount", True)
    s.value("tgt_mi", "Minority interest", t.minority_interest / amt, "amount", True)
    s.value("tgt_int", "Interest expense", t.interest_expense / amt, "amount", True)
    s.value("ppa_flag", "Book value entered? (1 = yes, runs the PPA)", 0 if t.book_value is None else 1, "number", True)
    s.value("tgt_book", "Book value of equity", (t.book_value or 0) / amt, "amount", True)

    o, st = d.offer, d.stake
    s.heading("Offer")
    s.value("price_input", 'Price input ("per_share" or "total")', o.price_input, "text", True)
    s.value("purchase_price", "Purchase price (when price input is total)", o.purchase_price / amt, "amount", True)
    s.value("offer_price", "Offer price per share", o.offer_price, "price", True)
    s.value("pct_stock", "Share of price paid in acquirer stock", o.pct_stock, "pct", True)
    s.value("debt_treatment", 'Target debt ("assumed" or "refinanced")', o.debt_treatment, "text", True)

    s.heading("Stake")
    s.value("existing", "Stake owned before the deal", st.existing_pct, "pct", True)
    s.value("bought", "Stake bought at the offer price", st.bought_pct, "pct", True)
    s.value("oo_pct", "Open offer: % of shares", st.open_offer_pct, "pct", True)
    s.value("oo_price", "Open offer price (0 = offer price)", st.open_offer_price, "price", True)
    s.value("acceptance", "Open offer acceptance", st.acceptance_pct, "pct", True)
    s.value("treatment_override", 'Accounting ("auto", "consolidate", "equity", "investment")', st.treatment, "text", True)
    s.value("inside", "Share of synergies realised inside the target", st.synergies_inside_target_pct, "pct", True)

    f = d.funding
    s.heading("Funding")
    s.value("cash_input", "Acquirer cash used", f.cash_used / amt, "amount", True)
    s.value("ne_amount", "New equity raised", f.new_equity_amount / amt, "amount", True)
    s.value("ne_price", "New equity issue price (0 = share price)", f.new_equity_price, "price", True)
    s.value("balancing", 'Balancing item ("none", "cash", "tranche_1"...)', f.balancing_item, "text", True)
    s.value("fcf", "Cash sweep: FCF conversion of cash net income", f.fcf_conversion_pct, "pct", True)
    for k in range(3):
        tr = f.tranches[k] if k < len(f.tranches) else None
        s.heading(f"Debt tranche {k + 1}")
        s.value(f"t{k}_name", "Name", tr.name if tr else "(unused)", "text", True)
        s.value(f"t{k}_amt", "Amount", (tr.amount if tr else 0) / amt, "amount", True)
        s.value(f"t{k}_rate", "Interest rate", tr.rate if tr else 0, "pct", True)
        s.value(f"t{k}_term", "Term", tr.term_years if tr else 5, "years", True)
        s.value(f"t{k}_amort", "Mandatory amortization per year (% of original)", tr.amortization_pct if tr else 0, "pct", True)
        s.value(f"t{k}_fee", "Financing fee", tr.fee_pct if tr else 0, "pct", True)

    s.heading("Fees and purchase price allocation")
    s.value("adv_pct", "Advisory fees (% of EV)", d.fees.advisory_pct, "pct", True)
    p = d.ppa
    s.value("int_mode", 'Intangibles input ("amount" or "pct_of_excess")', p.intangibles_input, "text", True)
    if p.intangibles_input == "amount":
        s.value("int_val", "Intangibles write-up", p.intangibles / amt, "amount", True)
    else:
        s.value("int_val", "Intangibles write-up (% of excess purchase price)", p.intangibles, "pct", True)
    s.value("int_life", "Intangibles useful life", p.intangibles_life, "years", True)
    s.value("ppe", "PP&E write-up", p.ppe_writeup / amt, "amount", True)
    s.value("ppe_life", "PP&E useful life", p.ppe_life, "years", True)

    y = d.synergies
    s.heading("Synergies", years=True)
    s.value("cost", "Cost synergies (full run-rate, pre-tax)", y.cost / amt, "amount", True)
    s.years("cost_ph", "Cost synergies phase-in", y.cost_phase_in, "pct", True)
    s.value("rev", "Revenue synergies (full run-rate)", y.revenue / amt, "amount", True)
    s.value("rev_margin", "Profit margin on revenue synergies", y.revenue_margin, "pct", True)
    s.years("rev_ph", "Revenue synergies phase-in", y.revenue_phase_in, "pct", True)
    s.value("prob", "Revenue synergies probability", y.revenue_probability, "pct", True)
    s.value("dis", "Dis-synergies (full, pre-tax)", y.dis_synergies / amt, "amount", True)
    s.years("dis_ph", "Dis-synergies phase-in", y.dis_synergy_phase_in, "pct", True)
    s.value("integ", "Integration costs (one-time total, pre-tax)", y.integration_costs / amt, "amount", True)
    s.years("integ_sch", "Integration costs spent", y.integration_schedule, "pct", True)

    s.heading("Credit thresholds and valuation")
    s.value("max_lev", "Maximum total debt / EBITDA", d.thresholds.max_leverage, "multiple", True)
    s.value("min_cov", "Minimum EBITDA / interest", d.thresholds.min_coverage, "multiple", True)
    s.value("disc", "Discount rate (r)", d.valuation.discount_rate, "pct", True)
    s.value("tg", "Terminal growth (g)", d.valuation.terminal_growth, "pct", True)


# ---------------------------------------------------------------- Sources & Uses

def write_sources_uses(book):
    s = book.sheet("Sources & Uses")
    s.title("Offer, consideration and sources & uses")

    s.heading("Offer (100% of the target)")
    s.value("ps", "Price per share? (1 = yes)", '=IF([price_input]="per_share",1,0)', "number")
    s.value("equity_value", "Implied equity value (100%)", "=IF([ps]=1,[offer_price]*[tgt_shares],[purchase_price])")
    s.value("net_debt", "Target net debt", "=[tgt_debt]-[tgt_cash]")
    s.value("ev", "Implied enterprise value", "=[equity_value]+[net_debt]+[tgt_pref]+[tgt_mi]")
    s.value("premium_pct", "Premium to unaffected price",
            '=IF(AND([ps]=1,[tgt_unaff]>0),[offer_price]/[tgt_unaff]-1,"n/a")', "pct")
    s.value("ev_rev", "EV / Revenue", '=IF([tgt_rev]>0,[ev]/[tgt_rev],"n/a")', "multiple")
    s.value("ev_ebitda", "EV / EBITDA", '=IF([tgt_ebitda]>0,[ev]/[tgt_ebitda],"n/a")', "multiple")
    s.value("pe", "Offer P/E",
            '=IF([tgt_ni]<=0,"n/a",IF([ps]=1,[offer_price]/([tgt_ni]/[tgt_shares]),[equity_value]/[tgt_ni]))',
            "multiple")

    s.heading("Stake")
    s.value("oo_price_eff", "Open offer price used", "=IF([oo_price]>0,[oo_price],[offer_price])", "price")
    s.value("neg_pct", "Stake bought (negotiated)", "=IF([ps]=1,[bought],1)", "pct")
    s.value("oo_eff_pct", "Stake bought (open offer, after acceptance)", "=IF([ps]=1,[oo_pct]*[acceptance],0)", "pct")
    s.value("existing_eff", "Stake owned before", "=IF([ps]=1,[existing],0)", "pct")
    s.value("final", "Final stake", "=[existing_eff]+[neg_pct]+[oo_eff_pct]", "pct")
    s.value("treatment", "Accounting treatment",
            '=IF([treatment_override]<>"auto",[treatment_override],'
            'IF([final]>0.5,"consolidate",IF([final]>=0.2,"equity","investment")))', "text")
    s.value("neg_shares", "Shares bought (negotiated)", "=IF([ps]=1,[neg_pct]*[tgt_shares],0)", "shares")
    s.value("oo_shares", "Shares bought (open offer)", "=[oo_eff_pct]*[tgt_shares]", "shares")
    s.value("eq_cost", "Cost of equity acquired",
            "=IF([ps]=1,[neg_shares]*[offer_price],[purchase_price])+[oo_shares]*[oo_price_eff]")
    s.value("premium_paid", "Premium paid over unaffected price",
            '=IF(AND([ps]=1,[tgt_unaff]>0),([offer_price]-[tgt_unaff])*[neg_shares]'
            '+([oo_price_eff]-[tgt_unaff])*[oo_shares],"n/a")')

    s.heading("Consideration")
    s.value("stock_value", "Paid in acquirer stock", "=[eq_cost]*[pct_stock]")
    s.value("exchange_ratio", "Exchange ratio (acquirer shares per target share)",
            '=IF([ps]=1,[offer_price]*[pct_stock]/[acq_price],"n/a")', "number")
    s.value("shares_issued", "Shares issued to target holders", "=[stock_value]/[acq_price]", "shares")
    s.value("ne_price_eff", "New equity issue price", "=IF([ne_price]>0,[ne_price],[acq_price])", "price")
    s.value("ne_shares", "Shares issued to new investors", "=[ne_amount]/[ne_price_eff]", "shares")

    s.heading("Uses of funds")
    s.value("use_equity", "Purchase of target equity", "=[eq_cost]")
    s.value("use_refi", "Target debt refinanced",
            '=IF(AND([debt_treatment]="refinanced",[treatment]="consolidate"),[tgt_debt],0)')
    s.value("use_adv", "Advisory fees", "=[adv_pct]*[ev]")
    s.value("cash_needed", "Cash needed before financing fees",
            "=[use_equity]+[use_refi]+[use_adv]-[stock_value]")

    s.heading("Sources of funds")
    s.value("src_stock", "Acquirer stock issued to target holders", "=[stock_value]")
    s.value("cash_eff", "Acquirer cash on hand",
            '=IF([balancing]="cash",MAX(0,[cash_needed]+[t0_amt]*[t0_fee]+[t1_amt]*[t1_fee]+[t2_amt]*[t2_fee]'
            '-[t0_amt]-[t1_amt]-[t2_amt]-[ne_amount]),[cash_input])')
    for k in range(3):
        others = [j for j in range(3) if j != k]
        other_amounts = "+".join(f"[t{j}_amt]" for j in others)
        other_fees = "+".join(f"[t{j}_amt]*[t{j}_fee]" for j in others)
        s.value(f"t{k}_eff", f"Debt tranche {k + 1}",
                f'=IF([balancing]="tranche_{k + 1}",MAX(0,([cash_needed]+{other_fees}'
                f'-[cash_input]-[ne_amount]-{other_amounts})/(1-[t{k}_fee])),[t{k}_amt])')
    s.value("src_ne", "New equity raised", "=[ne_amount]")
    s.value("use_fin_fees", "Financing fees (a use)", "=[t0_eff]*[t0_fee]+[t1_eff]*[t1_fee]+[t2_eff]*[t2_fee]")

    s.heading("Check")
    s.value("total_uses", "Total uses", "=[use_equity]+[use_refi]+[use_adv]+[use_fin_fees]")
    s.value("total_sources", "Total sources", "=[src_stock]+[cash_eff]+[t0_eff]+[t1_eff]+[t2_eff]+[src_ne]")
    s.value("gap", "Sources minus uses", "=[total_sources]-[total_uses]")
    s.value("balanced", "Status", f'=IF(ABS([gap])<={1 / book.units["amount"]},"Balanced","NOT BALANCED")', "text")
    s.value("cash_available", "Cash available above the minimum",
            '=IF([acq_cash_flag]=1,MAX(0,[acq_cash]-[acq_min_cash]),"n/a")')

    s.heading("Pro forma ownership")
    s.value("pf_shares", "Pro forma diluted shares", "=[acq_shares]+[shares_issued]+[ne_shares]", "shares")
    s.value("own_existing", "Existing acquirer holders", "=[acq_shares]/[pf_shares]", "pct")
    s.value("own_target", "Former target holders", "=[shares_issued]/[pf_shares]", "pct")
    s.value("own_new", "New investors", "=[ne_shares]/[pf_shares]", "pct")


# ---------------------------------------------------------------- PPA

def write_ppa(book):
    s = book.sheet("PPA")
    s.title("Purchase price allocation",
            "Partial goodwill: write-ups at 100%, goodwill on the acquirer's stake only. "
            "Applies only when the target is consolidated and a book value is entered.")
    s.value("ppa_on", "PPA applies? (1 = yes; not when the target was already controlled)",
            '=IF(AND([treatment]="consolidate",[ppa_flag]=1,[existing_eff]<=0.5),1,0)', "number")
    s.value("ppa_consideration", "Consideration (incl. prior stake at offer price)",
            "=[eq_cost]+IF([ps]=1,[existing_eff]*[tgt_shares]*[offer_price],0)")
    s.value("ppa_book", "Book value of equity", "=[tgt_book]")
    s.value("excess", "Excess purchase price (100% basis)", "=IF([final]>0,[ppa_consideration]/[final]-[ppa_book],0)")
    s.value("int_wu", "Intangibles write-up",
            '=[ppa_on]*IF([int_mode]="pct_of_excess",[int_val]*MAX([excess],0),[int_val])')
    s.value("ppe_wu", "PP&E write-up", "=[ppa_on]*[ppe]")
    s.value("dtl", "Deferred tax liability", "=([int_wu]+[ppe_wu])*[tax]")
    s.value("fvna", "Fair value of net assets", "=[ppa_book]+[int_wu]+[ppe_wu]-[dtl]")
    s.value("goodwill", "Goodwill", '=IF([ppa_on]=1,[ppa_consideration]-[final]*[fvna],"n/a")')
    s.heading("New amortization and depreciation (pre-tax)", years=True)
    s.years("amort", "Write-up amortization",
            lambda i: "=" + straight_line("[int_wu]", "[int_life]", i) + "+" + straight_line("[ppe_wu]", "[ppe_life]", i))


# ---------------------------------------------------------------- Debt

def write_debt(book):
    s = book.sheet("Debt Schedule")
    s.title("Debt schedule",
            "Interest is charged on the opening balance. The cash sweep (a simplification) uses "
            "pro forma cash net income x the FCF conversion %, after mandatory amortization, in tranche order.")
    s.heading("Cash for early repayment", years=True)
    s.years("mand_total", "Mandatory amortization, all tranches",
            lambda i: "=" + "+".join(f"[t{k}_mand@{i}]" for k in range(3)))
    s.years("sweep_pool", "Cash available for the sweep",
            lambda i: f"=MAX(0,MAX([cash_ni@{i}],0)*[fcf]-[mand_total@{i}])")

    for k in range(3):
        s.heading(f"Tranche {k + 1}", years=True)
        s.years(f"t{k}_beg", "Opening balance",
                lambda i: f"=[t{k}_eff]" if i == 0 else f"=[t{k}_end@{i - 1}]")
        s.years(f"t{k}_mand", "Mandatory amortization",
                lambda i: f"=MIN([t{k}_beg@{i}],[t{k}_amort]*[t{k}_eff])")
        earlier = "".join(f"-[t{j}_sweep@{{i}}]" for j in range(k))
        s.years(f"t{k}_sweep", "Cash sweep",
                lambda i: f"=MIN([t{k}_beg@{i}]-[t{k}_mand@{i}],[sweep_pool@{i}]{earlier.format(i=i)})")
        s.years(f"t{k}_end", "Closing balance",
                lambda i: f"=[t{k}_beg@{i}]-[t{k}_mand@{i}]-[t{k}_sweep@{i}]")
        s.years(f"t{k}_int", "Interest (on opening balance)", lambda i: f"=[t{k}_beg@{i}]*[t{k}_rate]")
        s.years(f"t{k}_feeam", "Financing fee amortization",
                lambda i: "=" + straight_line(f"[t{k}_eff]*[t{k}_fee]", f"[t{k}_term]", i))

    s.heading("Totals", years=True)
    s.years("debt_int", "Interest on new debt", lambda i: "=" + "+".join(f"[t{k}_int@{i}]" for k in range(3)))
    s.years("debt_fees", "Financing fee amortization", lambda i: "=" + "+".join(f"[t{k}_feeam@{i}]" for k in range(3)))
    s.years("debt_end", "Closing balance of new debt", lambda i: "=" + "+".join(f"[t{k}_end@{i}]" for k in range(3)))


# ---------------------------------------------------------------- Pro forma

def write_proforma(book):
    s = book.sheet("Pro Forma EPS")
    s.title("Pro forma EPS", "Accretion/dilution is measured against the acquirer's standalone EPS for the same year.")

    for prefix, name in (("acq", "Acquirer"), ("tgt", "Target")):
        s.heading(f"{name} standalone", years=True)
        s.years(f"{prefix}_rev_y", "Revenue",
                lambda i: f"=[{prefix}_rev]*(1+[{prefix}_rev_g@0])" if i == 0
                else f"=[{prefix}_rev_y@{i - 1}]*(1+[{prefix}_rev_g@{i}])")
        s.years(f"{prefix}_ebitda_y", "EBITDA",
                lambda i: f"=IF([{prefix}_rev]>0,[{prefix}_rev_y@{i}]*[{prefix}_margin@{i}],[{prefix}_ebitda])")
        s.years(f"{prefix}_ni_y", "Net income",
                lambda i: f"=[{prefix}_ni]*(1+[{prefix}_ni_g@0])" if i == 0
                else f"=[{prefix}_ni_y@{i - 1}]*(1+[{prefix}_ni_g@{i}])")
    s.note("Acquirer figures are as reported, so they already include any stake held in the target.")
    s.value("exist_share", "Existing stake's share of target income (20%+ stakes)",
            "=IF([existing_eff]>=0.2,[existing_eff],0)", "pct")
    s.years("exist_inc", "Of which: income from the existing stake", lambda i: f"=[exist_share]*[tgt_ni_y@{i}]")
    s.years("sa_eps", "Acquirer standalone EPS", lambda i: f"=[acq_ni_y@{i}]/[acq_shares]", "eps", bold=True)

    s.heading("Synergies (pre-tax, 100%)", years=True)
    s.years("syn_cost", "Cost synergies", lambda i: f"=[cost]*[cost_ph@{i}]")
    s.years("syn_rev", "Revenue synergy profit", lambda i: f"=[rev]*[rev_margin]*[rev_ph@{i}]*[prob]")
    s.years("syn_dis", "Dis-synergies", lambda i: f"=[dis]*[dis_ph@{i}]")
    s.years("syn_integ", "Integration costs", lambda i: f"=[integ]*[integ_sch@{i}]")
    s.years("syn_ebitda", "Synergies in EBITDA", lambda i: f"=[syn_cost@{i}]+[syn_rev@{i}]-[syn_dis@{i}]")

    s.heading("How the target is counted")
    s.value("cons", "Consolidated? (1 = yes)", '=IF([treatment]="consolidate",1,0)', "number")
    s.value("refi", "Target debt refinanced? (1 = yes)", '=IF(AND([cons]=1,[debt_treatment]="refinanced"),1,0)', "number")
    s.value("in_ni", "Share of synergies in net income",
            '=IF([cons]=1,1,IF([treatment]="equity",(1-[inside])+[final]*[inside],0))', "pct")
    s.value("attrib", "Share of synergies attributable to the acquirer",
            "=IF([cons]=1,1-(1-[final])*[inside],[in_ni])", "pct")
    s.value("in_ebitda", "Share of synergies in EBITDA",
            '=IF([cons]=1,1,IF([treatment]="equity",1-[inside],0))', "pct")

    s.heading("Pro forma net income (after tax)", years=True)
    s.years("c_acq", "Acquirer net income (as reported)", lambda i: f"=[acq_ni_y@{i}]")
    s.years("c_tgt", "Target net income (beyond the existing stake)",
            lambda i: f'=IF([cons]=1,[tgt_ni_y@{i}],IF([treatment]="equity",[final]*[tgt_ni_y@{i}],0))'
                      f"-[exist_inc@{i}]")
    s.years("c_refi", "Refinancing: target interest saved", lambda i: "=[refi]*[tgt_int]*(1-[tax])")
    s.years("c_cost", "Cost synergies", lambda i: f"=[syn_cost@{i}]*(1-[tax])*[in_ni]")
    s.years("c_rev", "Revenue synergies", lambda i: f"=[syn_rev@{i}]*(1-[tax])*[in_ni]")
    s.years("c_dis", "Dis-synergies", lambda i: f"=-[syn_dis@{i}]*(1-[tax])*[in_ni]")
    s.years("c_integ", "Integration costs", lambda i: f"=-[syn_integ@{i}]*(1-[tax])*[in_ni]")
    s.years("c_amort", "New amortization", lambda i: f"=-[cons]*[amort@{i}]*(1-[tax])")
    s.years("c_int", "Interest on new debt", lambda i: f"=-[debt_int@{i}]*(1-[tax])")
    s.years("c_fees", "Financing fee amortization", lambda i: f"=-[debt_fees@{i}]*(1-[tax])")
    s.years("c_lost", "Interest lost on cash used", lambda i: "=-[cash_eff]*[acq_cash_rate]*(1-[tax])")
    s.years("c_mi", "Minority interest",
            lambda i: f"=-[cons]*(1-[final])*([tgt_ni_y@{i}]+[inside]*([syn_ebitda@{i}]-[syn_integ@{i}])*(1-[tax])+[c_amort@{i}])")
    s.years("gaap_ni", "GAAP net income",
            lambda i: "=" + "+".join(f"[{k}@{i}]" for k in
                                     ("c_acq", "c_tgt", "c_refi", "c_cost", "c_rev", "c_dis", "c_integ",
                                      "c_amort", "c_int", "c_fees", "c_lost", "c_mi")), bold=True)

    s.heading("Earnings per share", years=True)
    s.years("pf_shares_y", "Pro forma diluted shares", lambda i: "=[pf_shares]", "shares")
    s.years("gaap_eps", "GAAP EPS", lambda i: f"=[gaap_ni@{i}]/[pf_shares_y@{i}]", "eps", bold=True)
    s.years("gaap_acc", "GAAP accretion / (dilution)", lambda i: f"=([gaap_eps@{i}]-[sa_eps@{i}])/ABS([sa_eps@{i}])", "pct", bold=True)
    s.years("cash_ni", "Cash net income (adds back amortization and integration costs)",
            lambda i: f"=[gaap_ni@{i}]-[c_amort@{i}]*[cons]*[final]+[syn_integ@{i}]*(1-[tax])*[attrib]")
    s.years("cash_eps", "Cash EPS", lambda i: f"=[cash_ni@{i}]/[pf_shares_y@{i}]", "eps", bold=True)
    s.years("cash_acc", "Cash accretion / (dilution)", lambda i: f"=([cash_eps@{i}]-[sa_eps@{i}])/ABS([sa_eps@{i}])", "pct", bold=True)
    s.years("breakeven", "Breakeven pre-tax synergies",
            lambda i: f'=IF([attrib]>0,([sa_eps@{i}]*[pf_shares_y@{i}]-([gaap_ni@{i}]-[syn_ebitda@{i}]*(1-[tax])*[attrib]))'
                      f'/((1-[tax])*[attrib]),"n/a")')
    s.note("Negative breakeven synergies mean the deal is accretive without any synergies.")
    s.value("one_time_fees", "One-time advisory fees (excluded from EPS)", "=[use_adv]")


# ---------------------------------------------------------------- Credit

def write_credit(book):
    s = book.sheet("Credit")
    s.title("Credit metrics", "End-of-year debt. Ratios need EBITDA figures to be entered.")
    s.value("ctrl_before", "Target already consolidated before the deal? (1 = yes)",
            "=IF([existing_eff]>0.5,1,0)", "number")
    s.value("adds_tgt", "Add the target's EBITDA, cash and debt? (1 = yes)",
            "=IF(AND([cons]=1,[ctrl_before]=0),1,0)", "number")
    s.value("has_ebitda", "EBITDA entered? (1 = yes)",
            "=IF(OR([acq_ebitda]>0,AND([adds_tgt]=1,[tgt_ebitda]>0)),1,0)", "number")
    s.heading("Combined company", years=True)
    s.years("cr_ebitda", "Pro forma EBITDA",
            lambda i: f"=[acq_ebitda_y@{i}]+[adds_tgt]*[tgt_ebitda_y@{i}]+[syn_ebitda@{i}]*[in_ebitda]")
    s.years("cr_debt", "Total debt",
            lambda i: f"=[acq_debt]+IF(AND([adds_tgt]=1,[refi]=0),[tgt_debt],0)"
                      f"-IF(AND([ctrl_before]=1,[refi]=1),[tgt_debt],0)+[debt_end@{i}]")
    s.years("cr_cash", "Cash", lambda i: "=[adds_tgt]*[tgt_cash]+IF([acq_cash_flag]=1,MAX([acq_cash]-[cash_eff],0),0)")
    s.years("cr_net_debt", "Net debt", lambda i: f"=[cr_debt@{i}]-[cr_cash@{i}]")
    s.years("cr_int", "Interest expense",
            lambda i: f"=[acq_int]+IF(AND([adds_tgt]=1,[refi]=0),[tgt_int],0)"
                      f"-IF(AND([ctrl_before]=1,[refi]=1),[tgt_int],0)+[debt_int@{i}]+[debt_fees@{i}]")
    ok = "AND([has_ebitda]=1,[cr_ebitda@{i}]>0)"
    s.heading("Ratios", years=True)
    s.years("cr_lev", "Total debt / EBITDA",
            lambda i: f'=IF({ok.format(i=i)},[cr_debt@{i}]/[cr_ebitda@{i}],"n/a")', "multiple")
    s.years("cr_net_lev", "Net debt / EBITDA",
            lambda i: f'=IF({ok.format(i=i)},[cr_net_debt@{i}]/[cr_ebitda@{i}],"n/a")', "multiple")
    s.years("cr_cov", "EBITDA / interest",
            lambda i: f'=IF(AND({ok.format(i=i)},[cr_int@{i}]>0),[cr_ebitda@{i}]/[cr_int@{i}],"n/a")', "multiple")
    s.years("cr_lev_flag", "Leverage check",
            lambda i: f'=IF({ok.format(i=i)},IF([cr_debt@{i}]/[cr_ebitda@{i}]>[max_lev],"BREACH","OK"),"n/a")', "text")
    s.years("cr_cov_flag", "Coverage check",
            lambda i: f'=IF(AND({ok.format(i=i)},[cr_int@{i}]>0),IF([cr_ebitda@{i}]/[cr_int@{i}]<[min_cov],"BREACH","OK"),"n/a")',
            "text")


# ---------------------------------------------------------------- Value

def write_value(book):
    s = book.sheet("Value Creation")
    s.title("Value creation", "End-of-year discounting. Terminal value = Year 5 run-rate x (1 + g) / (r - g).")
    s.heading("After-tax net synergies attributable to the acquirer", years=True)
    s.years("v_syn", "Net synergies", lambda i: f"=([syn_ebitda@{i}]-[syn_integ@{i}])*(1-[tax])*[attrib]")
    s.years("v_df", "Discount factor", lambda i: f"=1/(1+[disc])^{i + 1}", "number")
    s.years("v_pv", "Present value", lambda i: f"=[v_syn@{i}]*[v_df@{i}]")
    last = YEARS - 1
    s.heading("Summary")
    s.value("v_pv_years", "PV of Years 1-5", "=" + "+".join(f"[v_pv@{i}]" for i in range(YEARS)))
    s.value("v_run_rate", "Run-rate synergies (Year 5, excl. integration)",
            f"=[syn_ebitda@{last}]*(1-[tax])*[attrib]")
    s.value("v_tv", "Terminal value", '=IF([disc]>[tg],[v_run_rate]*(1+[tg])/([disc]-[tg]),"n/a")')
    s.value("v_pv_tv", "PV of terminal value", f'=IF([disc]>[tg],[v_tv]/(1+[disc])^{YEARS},"n/a")')
    s.value("v_total", "PV of synergies", '=IF([disc]>[tg],[v_pv_years]+[v_pv_tv],"n/a")')
    s.value("v_premium", "Premium paid", "=[premium_paid]")
    s.value("v_created", "Value created (PV of synergies - premium)",
            '=IF(AND([disc]>[tg],[ps]=1,[tgt_unaff]>0),[v_total]-[v_premium],"n/a")')


# ---------------------------------------------------------------- Sensitivities

def write_sensitivities(book, grid):
    s = book.sheet("Sensitivities")
    s.title("Sensitivities (fixed values)",
            "These grids are values calculated by the app, not formulas. Re-export after changing inputs.")
    if grid is None:
        s.note("No sensitivity grid was included in this export.")
        return
    x, y = grid["x_axis"], grid["y_axis"]
    kind_of = {"price": "pct", "premium": "pct", "synergies": "pct", "pct_stock": "pct",
               "debt_rate": "pct", "pct_debt": "pct", "acceptance": "pct"}
    for metric, name in (("gaap", "GAAP EPS accretion"), ("cash", "Cash EPS accretion")):
        for year in range(YEARS):
            s.heading(f"{name}, Year {year + 1}: {y['label']} (rows) vs {x['label']} (columns)")
            s.ws.cell(s.row, 1, f"{y['label']} \\ {x['label']}").font = Font(bold=True)
            for j, xv in enumerate(x["values"]):
                cell = s.ws.cell(s.row, FIRST_COL + j, xv)
                cell.number_format = FORMATS[kind_of[x["key"]]]
                cell.font = Font(bold=True)
            s.row += 1
            for r, yv in enumerate(y["values"]):
                cell = s.ws.cell(s.row, 1, yv)
                cell.number_format = FORMATS[kind_of[y["key"]]]
                cell.alignment = Alignment(horizontal="left")
                for j in range(len(x["values"])):
                    value = grid["cells"][r][j][metric][year]
                    c = s.ws.cell(s.row, FIRST_COL + j, value if value is not None else "n/a")
                    c.number_format = FORMATS["pct"]
                s.row += 1


def build_workbook(deal, grid=None):
    """Returns (workbook, refs). `refs` maps names to cell addresses (used by tests)."""
    d = apply_simple_mode(deal)
    book = Book(d.currency)
    write_inputs(book, d)
    write_sources_uses(book)
    write_ppa(book)
    write_debt(book)
    write_proforma(book)
    write_credit(book)
    write_value(book)
    write_sensitivities(book, grid)
    book.resolve()
    return book.wb, book.refs


def workbook_bytes(deal, grid=None):
    wb, _ = build_workbook(deal, grid)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
