# The Excel export: recalculate the workbook's live formulas (with pycel) and
# check they give the same numbers as the Python model, for several kinds of deal.
import io

import pytest
from openpyxl import load_workbook
from pycel import ExcelCompiler

from excel_export import build_workbook, workbook_bytes
from helpers import M, deal
from legacy_cases import CASE_2
from model.engine import run_model
from model.legacy import legacy_to_v2
from model.sensitivity import run_sensitivity

DEALS = {
    "simple mode (regression case 2)": legacy_to_v2(CASE_2),
    "advanced: balancing tranche, fees, PPA, growth, sweep": deal(
        offer={"pct_stock": 0.4, "debt_treatment": "refinanced"},
        acquirer={"revenue": 1000 * M, "ebitda": 200 * M, "debt": 300 * M, "cash": 150 * M, "min_cash": 20 * M,
                  "interest_expense": 15 * M, "interest_rate_on_cash": 0.03,
                  "net_income_growth": [0.05] * 5, "revenue_growth": [0.04] * 5},
        target={"revenue": 150 * M, "ebitda": 30 * M, "debt": 40 * M, "cash": 5 * M,
                "interest_expense": 3 * M, "book_value": 100 * M, "net_income_growth": [0.08] * 5},
        fees={"advisory_pct": 0.01},
        ppa={"intangibles_input": "pct_of_excess", "intangibles": 0.4, "intangibles_life": 8,
             "ppe_writeup": 10 * M, "ppe_life": 3},
        funding={"cash_used": 50 * M, "fcf_conversion_pct": 0.5, "balancing_item": "tranche_1",
                 "tranches": [{"rate": 0.07, "amortization_pct": 0.1, "fee_pct": 0.02, "term_years": 5},
                              {"name": "Notes", "amount": 60 * M, "rate": 0.08, "fee_pct": 0.015, "term_years": 7}]},
        synergies={"cost": 15 * M, "cost_phase_in": [0.5, 0.75, 1, 1, 1], "revenue": 40 * M, "revenue_margin": 0.2,
                   "revenue_phase_in": [0, 0.5, 1, 1, 1], "revenue_probability": 0.8, "dis_synergies": 2 * M,
                   "integration_costs": 30 * M, "integration_schedule": [0.6, 0.3, 0.1, 0, 0]},
        valuation={"discount_rate": 0.09, "terminal_growth": 0.02},
    ),
    "partial, consolidated, open offer, synergies inside target": deal(
        offer={"pct_stock": 0.0},
        stake={"existing_pct": 0.05, "bought_pct": 0.3, "open_offer_pct": 0.26, "open_offer_price": 27,
               "synergies_inside_target_pct": 0.5},
        acquirer={"ebitda": 200 * M},
        target={"ebitda": 30 * M, "debt": 40 * M, "book_value": 100 * M},
        ppa={"intangibles": 30 * M},
        funding={"balancing_item": "cash"},
        synergies={"cost": 15 * M, "integration_costs": 10 * M},
    ),
    "existing 25% stake bought up to 51%": deal(
        offer={"pct_stock": 0.5},
        stake={"existing_pct": 0.25, "bought_pct": 0.26},
        acquirer={"net_income": 105 * M, "ebitda": 200 * M},
        target={"ebitda": 30 * M, "book_value": 100 * M, "net_income_growth": [0.05] * 5},
        ppa={"intangibles": 20 * M},
        funding={"tranches": [{"rate": 0.06}], "balancing_item": "tranche_1"},
        synergies={"cost": 10 * M},
    ),
    "existing 60% stake, refinanced, bought to 100%": deal(
        offer={"pct_stock": 0.5, "debt_treatment": "refinanced"},
        stake={"existing_pct": 0.6, "bought_pct": 0.4},
        acquirer={"net_income": 112 * M, "ebitda": 230 * M, "debt": 540 * M, "interest_expense": 23 * M,
                  "cash": 80 * M},
        target={"ebitda": 30 * M, "debt": 40 * M, "cash": 5 * M, "interest_expense": 3 * M, "book_value": 100 * M},
        ppa={"intangibles": 20 * M},
        funding={"tranches": [{"rate": 0.06, "fee_pct": 0.01}], "balancing_item": "tranche_1"},
        synergies={"cost": 10 * M},
    ),
    "equity method (38%)": deal(
        offer={"pct_stock": 0.5},
        stake={"bought_pct": 0.25, "open_offer_pct": 0.26, "acceptance_pct": 0.5, "synergies_inside_target_pct": 0.4},
        funding={"tranches": [{"rate": 0.06}], "balancing_item": "tranche_1"},
        synergies={"cost": 10 * M},
    ),
}


def evaluator(d):
    wb, refs = build_workbook(d)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    compiler = ExcelCompiler(excel=load_workbook(buffer))
    return lambda ref: compiler.evaluate(ref.replace("$", "")), refs


@pytest.mark.parametrize("name", DEALS)
def test_excel_formulas_match_the_model(name):
    d = DEALS[name]
    expected = run_model(d)
    ev, refs = evaluator(d)
    scale = 1e6

    def amount(key):
        return ev(refs[key]) * scale

    su = expected["sources_uses"]
    assert amount("total_uses") == pytest.approx(su["total_uses"], rel=1e-9)
    assert amount("total_sources") == pytest.approx(su["total_sources"], rel=1e-9)
    assert ev(refs["pf_shares"]) * scale == pytest.approx(expected["ownership"]["pro_forma_shares"])
    assert ev(refs["treatment"]) == expected["stake"]["treatment"]
    if expected["ppa"]:
        assert amount("goodwill") == pytest.approx(expected["ppa"]["goodwill"])

    for i, year in enumerate(expected["years"]):
        assert ev(refs["gaap_eps"][i]) == pytest.approx(year["gaap_eps"], rel=1e-9), f"GAAP EPS year {i + 1}"
        assert ev(refs["cash_eps"][i]) == pytest.approx(year["cash_eps"], rel=1e-9), f"cash EPS year {i + 1}"
        assert ev(refs["gaap_acc"][i]) == pytest.approx(year["gaap_accretion"], rel=1e-9)
        assert ev(refs["c_mi"][i]) * scale == pytest.approx(year["components"]["minority_interest"], abs=1e-3)
        if year["breakeven_synergies"] is not None:
            assert ev(refs["breakeven"][i]) * scale == pytest.approx(year["breakeven_synergies"], rel=1e-9)
        credit = year["credit"]
        assert ev(refs["cr_debt"][i]) * scale == pytest.approx(credit["total_debt"], abs=1e-3)
        assert ev(refs["cr_ebitda"][i]) * scale == pytest.approx(credit["ebitda"], abs=1e-3)
        for key, field in (("cr_lev", "total_leverage"), ("cr_net_lev", "net_leverage"), ("cr_cov", "interest_coverage")):
            if credit[field] is None:
                assert ev(refs[key][i]) == "n/a"
            else:
                assert ev(refs[key][i]) == pytest.approx(credit[field])

    if expected["value"]["valid"]:
        assert amount("v_total") == pytest.approx(expected["value"]["pv_synergies"], rel=1e-9)


def test_inr_workbook_uses_crore_and_same_eps():
    d = deal(currency="INR", synergies={"cost": 15 * M})
    ev, refs = evaluator(d)
    assert ev(refs["acq_ni"]) == pytest.approx(10.0)  # 100M = 10 crore
    assert ev(refs["gaap_eps"][0]) == pytest.approx(run_model(d)["years"][0]["gaap_eps"])


def test_workbook_has_the_expected_tabs_and_sensitivities():
    d = DEALS["equity method (38%)"]
    data = workbook_bytes(d, run_sensitivity(d, "premium", "synergies"))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Inputs", "Sources & Uses", "PPA", "Debt Schedule", "Pro Forma EPS",
                             "Credit", "Value Creation", "Sensitivities"]
    values = [c.value for row in wb["Sensitivities"].iter_rows() for c in row if isinstance(c.value, float)]
    assert len(values) > 100
